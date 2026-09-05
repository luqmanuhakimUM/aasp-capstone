"""Source Organiser: resolves raw URLs/titles into formatted, annotated,
credibility-checked reference entries (FR-20..FR-25)."""
import io

from docx import Document
from docx.shared import Inches

from app.services import llm_client
from app.services.external_apis import resolve_single_entry

SYSTEM_PROMPT = """You write a single 1-2 sentence annotation summarizing what a paper is about and why it might be relevant to an undergraduate researcher, based only on its title/abstract metadata given. Do not evaluate credibility (that is handled separately). Be concise."""


def organise_sources(entries: list[str], language: str) -> list[dict]:
    results = []
    for raw_entry in entries:
        resolved = resolve_single_entry(raw_entry)
        if resolved is None:
            results.append(_unresolved_entry(raw_entry))
            continue
        results.append(_build_reference_entry(resolved, language))
    return results


def _build_reference_entry(source: dict, language: str) -> dict:
    tier, basis = _classify_credibility(source)
    annotation = _annotate(source, language)
    return {
        "title": source.get("title", "") or "Untitled",
        "authors": source.get("authors", ""),
        "year": source.get("year"),
        "doi": source.get("doi", ""),
        "url": source.get("url", ""),
        "credibility_tier": tier,
        "credibility_basis": basis,
        "annotation": annotation,
        "formatted_apa": format_apa(source),
        "resolved": True,
    }


def _unresolved_entry(raw_entry: str) -> dict:
    return {
        "title": raw_entry,
        "authors": "",
        "year": None,
        "doi": "",
        "url": raw_entry if raw_entry.lower().startswith("http") else "",
        "credibility_tier": "unknown",
        "credibility_basis": "Could not resolve this entry to a source record.",
        "annotation": "",
        "formatted_apa": "",
        "resolved": False,
    }


def _classify_credibility(source: dict) -> tuple[str, str]:
    api = source.get("source_api", "")
    venue = (source.get("venue") or "").strip()
    peer_reviewed_hint = source.get("peer_reviewed_hint", False)

    if api == "doaj":
        return "peer_reviewed", "Indexed in DOAJ, which only lists peer-reviewed open-access journals."
    if peer_reviewed_hint and venue:
        return "peer_reviewed", f"Published in '{venue}', identified as a peer-reviewed venue."
    if venue and not peer_reviewed_hint:
        return "grey_literature", f"Has a venue ('{venue}') but not confirmed as peer-reviewed (e.g. preprint, thesis, or report)."
    if api == "unresolved_url":
        return "non_academic", "No academic metadata found for this URL; treat as a non-peer-reviewed web source (blog/news/other) unless verified otherwise."
    return "unknown", "Insufficient metadata to determine credibility tier."


def _annotate(source: dict, language: str) -> str:
    if not source.get("abstract") and not source.get("title"):
        return ""
    user_prompt = f"""Title: {source.get('title', '')}
Abstract: {source.get('abstract', '') or '(not available)'}

{llm_client.language_directive(language)}"""
    try:
        return llm_client.complete_text(SYSTEM_PROMPT, user_prompt, max_tokens=150).strip()
    except Exception:
        return ""


def format_apa(source: dict) -> str:
    authors = source.get("authors", "").strip()
    year = source.get("year") or "n.d."
    title = source.get("title", "").strip().rstrip(".")
    venue = source.get("venue", "").strip()
    doi = source.get("doi", "").strip()

    author_str = authors if authors else "Author unknown"
    citation = f"{author_str} ({year}). {title}."
    if venue:
        citation += f" {venue}."
    if doi:
        citation += f" https://doi.org/{doi}"
    elif source.get("url"):
        citation += f" {source['url']}"
    return citation


def build_reference_docx(apa_citations: list[str], title: str = "References") -> bytes:
    """Renders a Word (.docx) reference list in standard APA 7 style:
    alphabetized, double-spaced, hanging indent."""
    doc = Document()
    doc.add_heading(title, level=1)
    for citation in sorted(c for c in apa_citations if c):
        paragraph = doc.add_paragraph(citation)
        fmt = paragraph.paragraph_format
        fmt.left_indent = Inches(0.5)
        fmt.first_line_indent = Inches(-0.5)
        fmt.line_spacing = 2
        fmt.space_after = 0
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
